from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _percent(value: float | None) -> str:
    return "Not used" if value is None else f"{value * 100:.1f}%"


def build_evidence_report(result: dict) -> bytes:
    """Create a compact Word report from one saved scan result."""

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)

    title = document.add_heading("Proofline Similarity Evidence Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(
        "Explainable passage-level evidence for responsible academic review"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    summary = document.add_table(rows=0, cols=2)
    summary.style = "Light Shading Accent 1"
    thresholds = result.get("review_thresholds", {})
    threshold_text = (
        f"show >= {_percent(thresholds['detection'])}; "
        f"medium >= {_percent(thresholds['medium'])}; "
        f"high >= {_percent(thresholds['high'])}"
        if {"detection", "medium", "high"}.issubset(thresholds)
        else "Not recorded"
    )
    summary_values = (
        ("Report ID", result.get("scan_id", "Not assigned")),
        ("Submitted document", result["submitted_document"]),
        ("Scan date (UTC)", result.get("created_at", "Not recorded")),
        ("Model mode", result["model_mode"]),
        ("Overall similarity", _percent(result["overall_score"])),
        ("Usable passages", str(result["total_passages"])),
        ("Evidence matches", str(result["matched_passages"])),
        ("Passages scored", str(result["reviewable_passages"])),
        ("Cited quotations excluded", str(result["excluded_passages"])),
        ("Review thresholds", threshold_text),
    )
    for label, value in summary_values:
        cells = summary.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)

    document.add_heading("Interpretation notice", level=1)
    document.add_paragraph(
        "This report presents textual-similarity evidence. It does not by itself prove "
        "plagiarism or academic misconduct; a qualified human reviewer must consider "
        "citations, quotations, assignment context, and institutional policy."
    )
    for warning in result.get("warnings", []):
        document.add_paragraph(warning, style="List Bullet")

    document.add_heading("Passage evidence", level=1)
    matches = result.get("matches", [])
    if not matches:
        document.add_paragraph("No passage exceeded the configured review threshold.")
    for index, match in enumerate(matches, start=1):
        document.add_heading(
            f"Match {index}: {match['review_band'].replace('_', ' ').title()}", level=2
        )
        document.add_paragraph(f"Source document: {match['source_document']}")
        passage_table = document.add_table(rows=2, cols=2)
        passage_table.style = "Table Grid"
        passage_table.cell(0, 0).text = "Submitted passage"
        passage_table.cell(0, 1).text = "Matching source passage"
        passage_table.cell(1, 0).text = match["submitted"]
        passage_table.cell(1, 1).text = match["source"]

        score_table = document.add_table(rows=1, cols=5)
        score_table.style = "Light Grid Accent 1"
        labels = ("Word", "Character", "Lexical", "Semantic", "Hybrid")
        keys = (
            "word_cosine_score",
            "character_jaccard_score",
            "lexical_score",
            "semantic_score",
            "hybrid_score",
        )
        for cell, label in zip(score_table.rows[0].cells, labels, strict=True):
            cell.text = label
        values = score_table.add_row().cells
        for cell, key in zip(values, keys, strict=True):
            cell.text = _percent(match.get(key))

        flags = []
        if match.get("is_quotation"):
            flags.append("quotation detected")
        if match.get("has_citation"):
            flags.append("citation detected")
        if match.get("excluded_from_overall"):
            flags.append("excluded from overall score")
        if flags:
            document.add_paragraph("Flags: " + ", ".join(flags))

    footer = section.footer.paragraphs[0]
    footer.text = "Generated locally by Proofline. Keep this report with the reviewed submission."
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
