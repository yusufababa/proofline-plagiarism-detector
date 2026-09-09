from __future__ import annotations

from pathlib import Path


class UnsupportedDocumentError(ValueError):
    pass


def extract_text(path: Path, allowed_extensions: frozenset[str]) -> str:
    suffix = path.suffix.lower()
    if suffix not in allowed_extensions:
        raise UnsupportedDocumentError(f"Unsupported document type: {suffix}")

    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Install python-docx to read DOCX files.") from exc
        document = Document(path)
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(blocks)

    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Install pypdf to read PDF files.") from exc
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)

